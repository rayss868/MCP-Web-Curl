from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

root = Path(__file__).resolve().parent / 'fixtures'
root.mkdir(parents=True, exist_ok=True)

# DOCX fixture
p = root / 'sample.docx'
doc = Document()
doc.add_heading('MCP-Web-Curl Thesis Fixture', level=1)
doc.add_paragraph('DOCX_MARKER_RAYHAN_2026')
doc.add_paragraph('This document is used only for deterministic parse_document testing.')
t = doc.add_table(rows=2, cols=2)
t.cell(0,0).text = 'Metric'; t.cell(0,1).text = 'Value'
t.cell(1,0).text = 'Expected'; t.cell(1,1).text = '42'
doc.save(p)

# PDF fixture
pdf_path = root / 'sample.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=A4)
c.setFont('Helvetica-Bold', 14)
c.drawString(72, 780, 'MCP-Web-Curl Thesis Fixture')
c.setFont('Helvetica', 11)
c.drawString(72, 750, 'PDF_MARKER_RAYHAN_2026')
c.drawString(72, 730, 'This PDF is used only for deterministic parse_document testing.')
c.save()

# Download fixture
(root / 'sample.txt').write_text('DOWNLOAD_MARKER_RAYHAN_2026\nDeterministic download fixture.\n', encoding='utf-8')
print(root)